#!/usr/bin/env python3
"""
generate_slides.py
==================

Builds ``final_defense.pptx`` - the B.Sc. Final-Defense presentation for
*Selective Intelligence* - on top of the department template
``Final-Defense Slide Template-PPT.pptx``.

The template supplies the slide masters, the two branded layouts (the DIU
logo, the top accent bar and the bottom defense bar) and the theme; this
script supplies every slide.  Nothing is a picture of text: titles, cards,
tables and charts are all native PowerPoint objects, so the deck stays
editable.  Screenshots of the running dashboard are the only bitmaps, and
they come from ``paper/figures/``.

Content is taken from ``final_paper.txt`` (report v7) - every number quoted
here appears in that report.

Run:  ../.venv/bin/python generate_slides.py
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "Final-Defense Slide Template-PPT.pptx"
OUTPUT = HERE / "final_defense.pptx"
FIGS = HERE / "figures"

# --------------------------------------------------------------- palette
INK = RGBColor(0x1F, 0x49, 0x7D)      # theme dk2 - titles
INK_D = RGBColor(0x15, 0x30, 0x54)    # deeper navy - dark cards
BLUE = RGBColor(0x4F, 0x81, 0xBD)     # theme accent1 - brand
BLUE_L = RGBColor(0x9D, 0xC3, 0xE6)
TEAL = RGBColor(0x21, 0x59, 0x68)     # template accent
TEAL_L = RGBColor(0x31, 0x86, 0x9B)
AMBER = RGBColor(0xC5, 0x5A, 0x11)    # findings / cautions
RED = RGBColor(0xB0, 0x3A, 0x2E)
GREEN = RGBColor(0x3F, 0x7D, 0x5A)
GREY = RGBColor(0x3F, 0x3F, 0x3F)     # body text
MUTED = RGBColor(0x76, 0x80, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF8, 0xFC)
CARD2 = RGBColor(0xEA, 0xF1, 0xF9)
BORDER = RGBColor(0xC9, 0xD8, 0xE8)

# ------------------------------------------------------------------ grid
L = 0.62            # left margin (in)
R = 12.71           # right edge
W = R - L           # 12.09
TOP = 1.28          # first content baseline
BOT = 6.34          # last content baseline (bottom bar starts at 6.50)

FONT = "Calibri"
MONO = "Consolas"


# =====================================================================
# low-level helpers
# =====================================================================
def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
         radius=None, lw=0.75):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    _no_shadow(sh)
    sh.text_frame.word_wrap = True
    return sh


def card(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.045):
    return rect(slide, x, y, w, h, fill=fill, line=line,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def frame_of(shape, pad_x=0.16, pad_y=0.10):
    """Turn a shape's own text frame into a padded text container."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad_x)
    tf.margin_top = tf.margin_bottom = Inches(pad_y)
    return tf


def _runs(p, text, size, color, bold=False, italic=False, font=FONT,
          strong_color=None):
    """Add runs to paragraph *p*, honouring **bold** segments."""
    strong_color = strong_color or color
    added = False
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if chunk == "":
            continue
        strong = chunk.startswith("**") and chunk.endswith("**")
        slant = (not strong) and chunk.startswith("*") and chunk.endswith("*")
        if strong:
            chunk = chunk[2:-2]
        elif slant:
            chunk = chunk[1:-1]
        r = p.add_run()
        r.text = chunk
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.italic = italic or slant
        f.bold = bold or strong
        f.color.rgb = strong_color if strong else color
        added = True
    if not added:                               # keep an empty spacer line
        r = p.add_run()
        r.text = " "
        r.font.size = Pt(size)
        r.font.name = font
    return p


def para(tf, text="", size=12, color=GREY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, before=0, after=3, bullet=None,
         bullet_color=None, indent=0.17, line=None, font=FONT,
         strong_color=None, first=False):
    if first:
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    _runs(p, text, size, color, bold=bold, italic=italic, font=font,
          strong_color=strong_color)
    if bullet:
        _set_bullet(p, bullet, bullet_color or BLUE, indent)
    return p


def _set_bullet(p, char, color, indent):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(int(Inches(indent)))))
    pPr.set("indent", str(-Emu(int(Inches(indent)))))
    from pptx.oxml.ns import nsmap
    from lxml import etree
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    srgb = etree.SubElement(buClr, qn("a:srgbClr"))
    srgb.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


def bullets(tf, items, size=12, color=GREY, char="▪", bullet_color=BLUE,
            after=6, indent=0.17, line=0.95):
    for it in items:
        para(tf, it, size=size, color=color, bullet=char,
             bullet_color=bullet_color, after=after, indent=indent, line=line)


# =====================================================================
# slide scaffolding
# =====================================================================
class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self._retitle_layouts()
        self._wipe_slides()
        self.title_layout = self.prs.slide_layouts[12]     # '1_Title Slide'
        self.body_layout = self.prs.slide_layouts[13]      # '1_Custom Layout'
        self._uid = 900

    # ---- template housekeeping ---------------------------------------
    def _retitle_layouts(self):
        """The stock template still says 'Pre-Defense' on the title layout."""
        for layout in self.prs.slide_layouts:
            for sh in layout.shapes:
                if not sh.has_text_frame:
                    continue
                txt = sh.text_frame.text.strip()
                if txt in ("B.Sc. Pre-Defense", "Final Year Defense"):
                    for p in sh.text_frame.paragraphs:
                        for i, r in enumerate(p.runs):
                            r.text = "B.Sc. Final-Defense" if i == 0 else ""
                            r.font.size = Pt(20 if "Title" in layout.name else 14)
                            r.font.bold = True

    def _wipe_slides(self):
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get(qn("r:id"))
            self.prs.part.drop_rel(rId)
            xml_slides.remove(sld)

    def next_id(self):
        self._uid += 1
        return self._uid

    # ---- slide factories ---------------------------------------------
    def _page_number(self, slide):
        for ph in slide.slide_layout.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                el = copy.deepcopy(ph._element)
                el.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set(
                    "id", str(self.next_id()))
                slide.shapes._spTree.append(el)
                sh = slide.shapes[-1]
                sh.left, sh.top = Inches(11.42), Inches(6.97)
                sh.width, sh.height = Inches(1.29), Inches(0.32)
                p = sh.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                p.font.size = Pt(11)
                p.font.color.rgb = MUTED
                p.font.name = FONT
                return

    def body(self, title, tag=None, rule=True):
        slide = self.prs.slides.add_slide(self.body_layout)
        t = slide.shapes.title
        t.left, t.top, t.width, t.height = (Inches(L), Inches(0.34),
                                            Inches(9.9), Inches(0.70))
        tf = t.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _runs(p, title, 26, INK, bold=True)
        if rule:
            rect(slide, L + 0.02, 1.07, 0.86, 0.072, fill=BLUE)
            rect(slide, L + 0.92, 1.07, 0.38, 0.072, fill=TEAL_L)
        if tag:
            chip = rect(slide, 10.34, 0.46, 2.37, 0.46, fill=CARD2,
                        line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        radius=0.28)
            ctf = frame_of(chip, 0.08, 0.02)
            ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(ctf, tag, size=11, color=INK, bold=True,
                 align=PP_ALIGN.CENTER, after=0, first=True)
        self._strip(slide)
        self._page_number(slide)
        return slide

    def _strip(self, slide):
        """Drop the inherited empty content placeholder."""
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type in (PP_PLACEHOLDER.OBJECT,
                                              PP_PLACEHOLDER.BODY) \
                    and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)

    def save(self):
        self.prs.save(str(OUTPUT))
        return OUTPUT


# =====================================================================
# composite components
# =====================================================================
def kpi(slide, x, y, w, h, value, label, fill=INK, vcolor=WHITE,
        lcolor=RGBColor(0xD6, 0xE4, 0xF2), vsize=26, lsize=10.5):
    sh = card(slide, x, y, w, h, fill=fill, line=fill)
    tf = frame_of(sh, 0.10, 0.09)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, value, size=vsize, color=vcolor, bold=True,
         align=PP_ALIGN.CENTER, after=1, first=True, line=0.9)
    para(tf, label, size=lsize, color=lcolor, align=PP_ALIGN.CENTER,
         after=0, line=0.92)
    return sh


def titled_card(slide, x, y, w, h, heading, items, accent=BLUE,
                fill=CARD, hsize=13, bsize=11.5, char="▪",
                body_color=GREY, note=None):
    sh = card(slide, x, y, w, h, fill=fill)
    rect(slide, x, y + 0.055, 0.075, h - 0.11, fill=accent)
    tf = frame_of(sh, 0.24, 0.11)
    para(tf, heading, size=hsize, color=INK, bold=True, after=5, first=True,
         line=0.95)
    for it in items:
        para(tf, it, size=bsize, color=body_color, bullet=char,
             bullet_color=accent, after=4, line=0.95)
    if note:
        para(tf, note, size=10, color=MUTED, italic=True, before=3, after=0,
             line=0.95)
    return sh


def banner(slide, x, y, w, h, text, fill=INK_D, color=WHITE, size=12.5,
           label=None, label_color=BLUE_L, align=PP_ALIGN.LEFT):
    sh = card(slide, x, y, w, h, fill=fill, line=fill)
    tf = frame_of(sh, 0.22, 0.10)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if label:
        para(tf, label, size=9.5, color=label_color, bold=True, after=2,
             first=True, align=align)
        para(tf, text, size=size, color=color, after=0, align=align,
             line=0.98, strong_color=WHITE)
    else:
        para(tf, text, size=size, color=color, after=0, first=True,
             align=align, line=0.98, strong_color=WHITE)
    return sh


def steps(slide, x, y, w, h, labels, fill=BLUE, color=WHITE, size=11.5,
          gap=0.06):
    n = len(labels)
    seg = (w - gap * (n - 1)) / n
    shapes = []
    for i, lab in enumerate(labels):
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = rect(slide, x + i * (seg + gap), y, seg, h,
                  fill=fill, line=fill, shape=shp)
        try:
            sh.adjustments[0] = 0.22
        except Exception:
            pass
        tf = frame_of(sh, 0.05, 0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, lab, size=size, color=color, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True, line=0.92)
        shapes.append(sh)
    return shapes


def table(slide, x, y, w, data, widths=None, header_fill=INK,
          font=10, hfont=10, row_h=0.30, head_h=0.34, aligns=None,
          zebra=CARD, first_col_bold=False, bold_rows=(), col_colors=None):
    rows, cols = len(data), len(data[0])
    h = head_h + row_h * (rows - 1)
    gf = slide.shapes.add_table(rows, cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    if widths:
        total = sum(widths)
        for i, cw in enumerate(widths):
            tbl.columns[i].width = Emu(int(Inches(w * cw / total)))
    tbl.rows[0].height = Emu(int(Inches(head_h)))
    for r in range(1, rows):
        tbl.rows[r].height = Emu(int(Inches(row_h)))
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            elif zebra is not None and r % 2 == 0:
                cell.fill.fore_color.rgb = zebra
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            al = PP_ALIGN.LEFT
            if aligns and aligns[c] == "c":
                al = PP_ALIGN.CENTER
            elif aligns and aligns[c] == "r":
                al = PP_ALIGN.RIGHT
            p.alignment = al
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = 0.94
            if r == 0:
                _runs(p, str(val), hfont, WHITE, bold=True)
            else:
                col = GREY
                if col_colors and col_colors.get(c):
                    col = col_colors[c]
                bold = (first_col_bold and c == 0) or (r in bold_rows)
                _runs(p, str(val), font, col, bold=bold, strong_color=INK)
    return gf


def shot(slide, name, x, y, w=None, h=None, caption=None, cap_size=9.5,
         border=True):
    path = FIGS / name
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    if border:
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(1)
    if caption:
        tf = textbox(slide, x, y + Emu(pic.height).inches + 0.055,
                     Emu(pic.width).inches, 0.3)
        para(tf, caption, size=cap_size, color=MUTED, align=PP_ALIGN.CENTER,
             after=0, first=True, line=0.95)
    return pic


# ------------------------------------------------------------- charts
def _style_axes(chart, value_axis=False, cat_size=10.5):
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = BORDER
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.color.rgb = GREY
    ca.tick_labels.font.name = FONT
    va = chart.value_axis
    va.has_major_gridlines = False
    va.visible = value_axis
    if value_axis:
        va.tick_labels.font.size = Pt(9.5)
        va.tick_labels.font.color.rgb = MUTED


def bar_chart(slide, x, y, w, h, cats, series, colors, num_fmt='0',
              legend=False, gap=55, overlap=-18, label_size=10.5,
              label_pos=XL_LABEL_POSITION.OUTSIDE_END, horizontal=False,
              y_max=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ctype = XL_CHART_TYPE.BAR_CLUSTERED if horizontal \
        else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(ctype, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    chart = gf.chart
    chart.font.size = Pt(10.5)
    chart.font.name = FONT
    chart.font.color.rgb = GREY
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10.5)
    plot = chart.plots[0]
    plot.gap_width = gap
    if len(series) > 1:
        plot.overlap = overlap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = num_fmt
    dl.number_format_is_linked = False
    dl.position = label_pos
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.color.rgb = INK
    for s, c in zip(chart.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c
        s.format.line.fill.background()
    if y_max is not None:
        chart.value_axis.maximum_scale = y_max
    _style_axes(chart)
    return chart


def pie_chart(slide, x, y, w, h, cats, vals, colors, num_fmt='0.0%',
              show_pct=True, legend=True, label_size=10.5, labels=True):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("s", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    chart = gf.chart
    chart.font.size = Pt(10.5)
    chart.font.name = FONT
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.color.rgb = GREY
    plot = chart.plots[0]
    plot.has_data_labels = labels
    if labels:
        dl = plot.data_labels
        dl.show_percentage = show_pct
        dl.show_value = not show_pct
        dl.number_format = num_fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(label_size)
        dl.font.bold = True
        dl.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    pts = chart.series[0].points
    for pt, c in zip(pts, colors):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
        pt.format.line.color.rgb = WHITE
        pt.format.line.width = Pt(1.5)
    return chart


# =====================================================================
# the deck
# =====================================================================
def build():
    d = Deck()

    # ---------------------------------------------------------- 1 title
    s = d.prs.slides.add_slide(d.title_layout)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = (Inches(0.75), Inches(1.02),
                                        Inches(11.83), Inches(2.05))
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(4)
    p.line_spacing = 0.92
    _runs(p, "Selective Intelligence", 40, INK, bold=True)
    para(tf, "A Cost-Aware Hybrid NLP–LLM Microservice for Thread-Level "
             "Sentiment and Watchlist-Driven Target Stance over Code-Mixed "
             "Bangla Social Media, with an MCP-Backed Agentic RAG Insight "
             "Layer", size=14.5, color=TEAL, align=PP_ALIGN.CENTER,
         after=0, line=1.0)

    sub = s.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = (Inches(3.39), Inches(4.26),
                                                Inches(6.55), Inches(1.30))
    stf = sub.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = 0
    para(stf, "Bishwajit Kumar Chakraborty", size=19, color=RGBColor(0, 0, 0),
         bold=True, align=PP_ALIGN.CENTER, after=2, first=True,
         font="Times New Roman")
    para(stf, "ID : 0242220005101414", size=15, color=GREY,
         align=PP_ALIGN.CENTER, after=2, font="Times New Roman")
    para(stf, "Department of CSE  ·  Daffodil International University",
         size=13.5, color=GREY, align=PP_ALIGN.CENTER, after=0,
         font="Times New Roman")

    sup = card(s, 1.55, 5.62, 5.05, 0.96, fill=CARD, line=BORDER)
    tfx = frame_of(sup, 0.16, 0.09)
    para(tfx, "SUPERVISED BY", size=9, color=BLUE, bold=True, after=2,
         first=True)
    para(tfx, "Dr. Sheak Rashed Haider Noori", size=12.5, color=INK,
         bold=True, after=1)
    para(tfx, "Professor & Head, Department of CSE, DIU", size=10.5,
         color=GREY, after=0)

    cosup = card(s, 6.75, 5.62, 5.05, 0.96, fill=CARD, line=BORDER)
    tfx = frame_of(cosup, 0.16, 0.09)
    para(tfx, "CO-SUPERVISED BY", size=9, color=TEAL_L, bold=True, after=2,
         first=True)
    para(tfx, "Dr. Md Alamgir Kabir", size=12.5, color=INK, bold=True,
         after=1)
    para(tfx, "Assistant Professor, Department of CSE, DIU", size=10.5,
         color=GREY, after=0)

    dt = textbox(s, 0.62, 3.16, 12.09, 0.3)
    para(dt, "B.Sc. in Computer Science and Engineering  ·  Final Defense  ·  "
             "August 2026", size=11.5, color=MUTED, align=PP_ALIGN.CENTER,
         after=0, first=True)
    rect(s, 0.0, 7.02, 13.333, 0.075, fill=BLUE)
    ft = textbox(s, 0.62, 6.70, 12.09, 0.28)
    para(ft, "Department of Computer Science and Engineering  ·  Faculty of "
             "Science and Information Technology  ·  Dhaka, Bangladesh",
         size=10.5, color=MUTED, align=PP_ALIGN.CENTER, after=0, first=True)
    d._page_number(s)

    # -------------------------------------------------------- 2 outline
    s = d.body("Outline", tag="Roadmap")
    items = [
        ("01", "Introduction", "the text, the platform, the bill"),
        ("02", "Motivation & Problem Statement", "a cost asymmetry worth designing for"),
        ("03", "Objectives", "O1 – O4, each checkable against the running system"),
        ("04", "Background Study", "Bangla & code-mixed NLP, cascades, RAG, agents"),
        ("05", "Gap Analysis", "four gaps, four objectives, one for one"),
        ("06", "Research Methodology", "five stages, six gates, eight labellers"),
        ("07", "Results & Analysis", "routing, cost split, retrieval, provenance"),
        ("08", "Novelty of the Work", "what is new, and what measurement disconfirmed"),
        ("09", "Sample Dataset & Expected Output", "50 posts, 10,272 comments, one JSON"),
        ("10", "Web Interface", "eleven tabs over a 51-path API"),
        ("11", "Demonstration Video", "the system running end to end"),
        ("12", "Conclusion, Limitations & References", "what is claimed, and where it stops"),
    ]
    cw, ch, gx, gy = 5.92, 0.72, 0.25, 0.115
    for i, (num, head, note) in enumerate(items):
        col, row = i % 2, i // 2
        x = L + col * (cw + gx)
        y = TOP + 0.04 + row * (ch + gy)
        sh = card(s, x, y, cw, ch, fill=CARD if row % 2 == 0 else CARD2)
        badge = rect(s, x + 0.09, y + 0.09, 0.54, ch - 0.18, fill=INK,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
        btf = frame_of(badge, 0.02, 0.02)
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(btf, num, size=14, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True)
        tf = textbox(s, x + 0.76, y + 0.10, cw - 0.9, ch - 0.2)
        para(tf, head, size=13, color=INK, bold=True, after=1, first=True,
             line=0.95)
        para(tf, note, size=10, color=MUTED, after=0, line=0.95)

    # --------------------------------------------------- 3 introduction
    s = d.body("Introduction", tag="Chapter 1")
    titled_card(s, L, TOP, 3.87, 2.62, "The text is not one language",
                ["Bangla script, English and romanized **Banglish** mix "
                 "inside single comments, with no standard orthography.",
                 "Measured over 10,272 comments: **81.4 %** Bangla script "
                 "only, **11.1 %** Latin only, **4.6 %** both, **2.9 %** "
                 "neither (pure emoji).",
                 "One comment in six is where a Bangla-only and an "
                 "English-only model each read the wrong language."],
                accent=BLUE)
    titled_card(s, L + 4.11, TOP, 3.87, 2.62,
                "The upstream platform does not understand it",
                ["It scrapes reliably and keys records by stable identifiers "
                 "— then stores **one coarse sentiment per post**.",
                 "Per-comment sentiment is **null on all 10,272 comments**; "
                 "no OCR; no embeddings.",
                 "It is a **data source, not an analysis layer** — which is "
                 "the gap this microservice fills."],
                accent=TEAL_L)
    titled_card(s, L + 8.22, TOP, 3.87, 2.62,
                "Both obvious answers fail",
                ["**Send everything to an LLM:** accurate, unaffordable — "
                 "cost scales with comments, not posts, and one viral thread "
                 "here carries **2,857** comments.",
                 "**Run only cheap classifiers:** affordable, and "
                 "demonstrably weaker on exactly the code-mixed register "
                 "that dominates the traffic."],
                accent=AMBER)
    banner(s, L, 4.10, W, 1.10,
           "Given a post together with its embedded comment thread, decide "
           "**per unit of work** how much intelligence that unit deserves, "
           "spend the expensive model only where cheap models cannot answer, "
           "and prove afterwards — from the system's own counters — where "
           "the money went and what produced every label.",
           label="PROBLEM STATEMENT", size=13)
    for i, (v, lab) in enumerate([("50", "posts · 25 campaigns"),
                                  ("10,272", "comments analysed"),
                                  ("3.75 %", "of comments the platform holds"),
                                  ("0", "comments labelled upstream")]):
        kpi(s, L + i * 3.055, 5.42, 2.86, 0.86, v, lab,
            fill=TEAL if i % 2 else INK, vsize=22)

    # ---------------------------------------------------- 4 motivation
    s = d.body("Motivation: a cost asymmetry worth designing for",
               tag="Chapter 1")
    c = card(s, L, TOP, 5.85, 1.72, fill=CARD)
    rect(s, L, TOP + 0.055, 0.075, 1.61, fill=GREEN)
    tf = frame_of(c, 0.24, 0.11)
    para(tf, "Cheap path — 7 small classifier heads", size=13, color=INK,
         bold=True, after=4, first=True)
    bullets(tf, ["135 – 280 M parameters, batched on **CPU**",
                 "**~100 ms** per short comment, **zero** per-token cost",
                 "Weak precisely on the romanized register"],
            size=11.5, bullet_color=GREEN, after=3)
    c = card(s, L + 6.24, TOP, 5.85, 1.72, fill=CARD)
    rect(s, L + 6.24, TOP + 0.055, 0.075, 1.61, fill=RED)
    tf = frame_of(c, 0.24, 0.11)
    para(tf, "Expensive path — one 7 B instruction-tuned model", size=13,
         color=INK, bold=True, after=4, first=True)
    bullets(tf, ["**7 – 50 s** for a single post-level call, locally served",
                 "A Stage-2 batch of 25 comments is a **generation** call — "
                 "cost scales with prompt length",
                 "Better on code-mixed text, and it is the only labeller "
                 "that sees the parent post"],
            size=11.5, bullet_color=RED, after=3)
    banner(s, L, 3.16, W, 0.62,
           "**Two to three orders of magnitude** in latency, and a "
           "comparable gap in money on a metered backend. A system that "
           "decides *which* items need the expensive path converts an "
           "unbounded bill into a bounded one.", fill=INK, size=12.5)
    titled_card(s, L, 3.94, 5.85, 1.36, "Why the decision is hard to make well",
                ["A confidence gate is only as good as the estimate feeding "
                 "it — and an uncalibrated estimate on an unseen language is "
                 "**a guess wearing a number**.",
                 "FrugalGPT [30], RouteLLM [31] and Hybrid LLM [32] fix the "
                 "shape of the answer **for English single queries**."],
                accent=AMBER)
    titled_card(s, L + 6.24, 3.94, 5.85, 1.36,
                "Why sentiment alone is not the question",
                ["Monitoring needs the **target** of the affect, not only "
                 "its polarity.",
                 "The hard part is the **matcher**: one entity is written in "
                 "three scripts, and a watchlist that matches one spelling "
                 "**fails silently** — zero mentions looks exactly like an "
                 "entity nobody discussed."],
                accent=TEAL_L)
    tfn = textbox(s, L, 5.48, W, 0.7)
    para(tfn, "No published cascade routes a *container* whose thousands of "
              "sub-items each need labelling — so none reports what share of "
              "the bill the routing decision actually governs. That "
              "unreported quantity is where this project starts.",
         size=12, color=INK, italic=True, after=0, first=True, line=1.0)

    # ---------------------------------------------------- 5 objectives
    s = d.body("Objectives", tag="Chapter 1 · § 1.3")
    obj = [
        ("O1", "Build the analysis service", "closes G1",
         "A horizontally scalable, queue-based microservice that consumes a "
         "post-with-details payload, analyses it in its **own** datastores, "
         "and returns one schema-validated JSON document per thread — "
         "**never writing back** to the upstream platform.", BLUE),
        ("O2", "Make the cascade selective — and measure it", "closes G2",
         "Cheap NLP on every post and every comment; a rule-based router "
         "decides post-level spend and, **separately**, which comments the "
         "expensive labellers read. Then report the routing rate, the "
         "post/comment call split and the **share of spend the gate governs** "
         "from the system's own counters.", TEAL),
        ("O3", "Deliver the capability layers the setting needs", "closes G3",
         "An eight-labeller per-comment ensemble in which a labeller that "
         "cannot answer **abstains**; alias-aware watchlist target stance "
         "kept in a separate field from sentiment; and an agentic insight "
         "layer over MCP tools under budgets, injection hardening and "
         "citation verification.", AMBER),
        ("O4", "Make every claim auditable", "closes G4",
         "Provenance on every label, per-lane cost telemetry, a live "
         "per-post trace and a regression suite that fails when a repaired "
         "defect returns — with an **evidence class attached to every "
         "published claim**, so the boundary of the evaluation is part of "
         "the result.", GREEN),
    ]
    y = TOP + 0.02
    for code, head, gap, body, accent in obj:
        h = 1.16
        sh = card(s, L, y, W, h, fill=CARD)
        rect(s, L, y + 0.055, 0.075, h - 0.11, fill=accent)
        badge = rect(s, L + 0.22, y + 0.24, 0.72, 0.68, fill=accent,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        btf = frame_of(badge, 0.02, 0.02)
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(btf, code, size=18, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True)
        tf = textbox(s, L + 1.06, y + 0.13, 9.55, h - 0.26)
        para(tf, head, size=13.5, color=INK, bold=True, after=2, first=True)
        para(tf, body, size=11.5, color=GREY, after=0, line=0.97)
        tag = rect(s, R - 1.28, y + 0.13, 1.06, 0.30, fill=WHITE,
                   line=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                   radius=0.30)
        ttf = frame_of(tag, 0.04, 0.01)
        ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ttf, gap, size=9.5, color=accent, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True)
        y += h + 0.13

    # --------------------------------------------- 6 background study I
    s = d.body("Background Study — Bangla, code-mixing and stance",
               tag="Chapter 2 · § 2.2")
    data = [["Author(s)", "Year", "Contribution", "Why it matters here"],
            ["Bhattacharjee et al. [1]", "2022",
             "BanglaBERT — ELECTRA-style pretraining on 27.5 GB of Bangla",
             "Strongest monolingual Bangla encoder; **two of our seven "
             "ensemble heads** are BanglaBERT derivatives"],
            ["Islam et al. [2]", "2021",
             "SentNoB — 15,000 human-annotated noisy Bangla comments",
             "Social-media Bangla is materially harder than curated Bangla; "
             "the SentNoB fine-tune is one of our voters"],
            ["Sazzed [4]", "2021",
             "Abusive content in a transliterated Bengali–English corpus",
             "Exactly our setting: **transliteration, not vocabulary**, is "
             "the primary source of error"],
            ["Nezami & Huda [6]", "2024",
             "Sentiment for code-mixed English–Bangla and Banglish text",
             "Closest published formulation to our comment lane; Banglish is "
             "a **distinct third bucket**"],
            ["Alam et al. [9]", "2021",
             "Survey of Bangla NLP tasks and transformer utility",
             "**Annotated data, not architecture, is the bottleneck** — why "
             "our effort goes into the system and its provenance"],
            ["Kucuk & Can [13]", "2021", "Stance detection: a survey",
             "Stance toward a target and document sentiment are **distinct "
             "quantities** — hence two separate output fields"],
            ["Patwa et al. [10]", "2020",
             "SemEval-2020 Task 9 — code-mixed sentiment",
             "Fixes the standard evaluation protocol, and the size of the "
             "gap versus monolingual data"]]
    table(s, L, TOP + 0.02, W, data, widths=[2.0, 0.62, 3.5, 5.1],
          font=10, hfont=10.5, row_h=0.545, head_h=0.34,
          aligns=[None, "c", None, None], first_col_bold=True)
    tfn = textbox(s, L, 5.72, W, 0.5)
    para(tfn, "Takeaway — the field supplies good models and repeatedly says "
              "the same thing: **romanized text is where they fail, and "
              "labels are the binding constraint.**", size=12, color=INK,
         after=0, first=True, line=1.0)

    # -------------------------------------------- 7 background study II
    s = d.body("Background Study — cascades, retrieval and agents",
               tag="Chapter 2 · § 2.2")
    data = [["Author(s)", "Year", "Contribution", "Why it matters here"],
            ["Chen, Zaharia & Zou [30]", "2023",
             "FrugalGPT — LLM cascade with a learned scorer and a budget",
             "The foundational cost-cascade result; we implement the idea "
             "with a **rule-based** gate and measure where the analogy "
             "breaks for thread-shaped units"],
            ["Ong et al. [31]", "2024",
             "RouteLLM — learning to route from preference data",
             "The routing decision **can be learned**; named in Chapter 6 as "
             "the most valuable single upgrade to our hand-tuned gate"],
            ["Lewis et al. [34]; Gao et al. [35]", "2020 · 2023",
             "Retrieval-augmented generation, and its survey",
             "The pattern our insight layer generalises from single-shot "
             "retrieval to a **tool-using loop**"],
            ["Hou et al. [38]", "2025",
             "Model Context Protocol — landscape and security threats",
             "The tool-interface standard we adopt, together with the "
             "threats that motivated the runner's hardening"],
            ["Greshake et al. [43]", "2023",
             "Indirect prompt injection in LLM-integrated applications",
             "Retrieved third-party comment text is **adversarial input**; "
             "hardening is risk reduction, never a solution"]]
    table(s, L, TOP + 0.02, W, data, widths=[2.15, 0.75, 3.4, 4.9],
          font=10, hfont=10.5, row_h=0.52, head_h=0.34,
          aligns=[None, "c", None, None], first_col_bold=True)
    y = 4.86
    tfh = textbox(s, L, y - 0.30, W, 0.3)
    para(tfh, "Five classes of similar application — and how each one falls "
              "short", size=13, color=INK, bold=True, after=0, first=True)
    sims = [("Commercial listening suites",
             "Complete but closed, English-centric, comments sampled"),
            ("The upstream platform",
             "Reliable ingestion; comment sentiment null, no OCR"),
            ("Academic Bangla classifiers",
             "Good models, no system: no ingestion, serving or cost model"),
            ("General RAG assistants",
             "Retrieve text a user could already read; no injection defence"),
            ("Cascade / router frameworks",
             "Components, not systems: no ingestion contract, no telemetry")]
    cwid = (W - 4 * 0.16) / 5
    for i, (h1, h2) in enumerate(sims):
        x = L + i * (cwid + 0.16)
        sh = card(s, x, y, cwid, 1.20, fill=CARD2)
        tf = frame_of(sh, 0.13, 0.10)
        para(tf, h1, size=11.5, color=INK, bold=True, after=3, first=True,
             line=0.95)
        para(tf, h2, size=10, color=GREY, after=0, line=0.95)

    # ------------------------------------------------- 8 gap analysis
    s = d.body("Gap Analysis", tag="Chapter 2 · § 2.3")
    data = [["Gap", "Observed in prior work", "Obj.", "How this project closes it"],
            ["**G1** — No deployable analysis layer for code-mixed Bangla threads",
             "The capabilities sit in classes of system that never meet: "
             "platforms ingest and barely analyse; listening suites are "
             "closed and English-centric; academic work ships classifiers, "
             "not systems.", "O1",
             "One queue-based microservice: ingest the whole stored thread, "
             "analyse in its own stores, emit one schema-validated JSON per "
             "thread, never write back — the platform's coarse score kept "
             "**beside** the recomputed one."],
            ["**G2** — Cascades are evaluated on single queries, so the cost "
             "of routing a *container* is unreported",
             "The cost-aware cascade literature [30], [31], [32] routes one "
             "query to one model. No published cascade routes a thread whose "
             "sub-items each need labelling, so none reports what share of "
             "the bill the container-level decision governs.", "O2",
             "The router makes **two separate decisions**, and the counters "
             "report the post- versus comment-level call split. Measured "
             "answer: the gate governs only **30 %** of calls — a finding "
             "that contradicted this project's own initial claim."],
            ["**G3** — The capability layers an adversarial code-mixed "
             "setting needs are missing or unsafe",
             "Fused ensembles hide a member that failed to load; "
             "target stance assumes English and a fixed target set; agentic "
             "RAG is judged on whether the answer *looks* right, with no "
             "defence against instructions inside retrieved text [43].",
             "O3",
             "Eight labellers with **explicit abstention** and *uncertain* "
             "where nothing read the comment; a versioned, alias-aware "
             "watchlist built for the three-script problem; and a runner "
             "with untrusted-data delimiters, citation verification and "
             "non-answer rejection."],
            ["**G4** — Systems papers report what a system does, not which "
             "claims were checked",
             "An implemented-but-unexercised capability reads exactly like a "
             "measured one, and a green test suite reads like a quality "
             "claim. Where the evidence stops is left to the reader.", "O4",
             "Every capability carries an **evidence class** — measured, "
             "works-but-unmeasured, unexercised, out of scope — recorded "
             "once. The image modality is reported as unexercised; labelling "
             "accuracy as a stated scope boundary."]]
    table(s, L, TOP + 0.02, W, data, widths=[2.55, 3.55, 0.62, 4.05],
          font=9.5, hfont=10.5, row_h=1.09, head_h=0.32,
          aligns=[None, None, "c", None])
    tfn = textbox(s, L, 6.02, W, 0.32)
    para(tfn, "Four gaps, four objectives, **one for one** — the gap analysis "
              "is the specification for Chapter 3.", size=12, color=INK,
         after=0, first=True)

    # ------------------------------------------------- 9 methodology
    s = d.body("Research Methodology — how the work was carried out",
               tag="Chapter 3")
    steps(s, L, TOP + 0.04, W, 0.62,
          ["Requirement\nAnalysis", "System\nDesign", "Development",
           "Testing", "Deployment", "Evaluation\n& Audit"])
    rows = [("Requirement Analysis",
             "Read the upstream contract for what it does and does **not** "
             "supply; write functional and non-functional requirements "
             "against it."),
            ("System Design",
             "Five decoupled stages over Redis Streams; the router's two "
             "decisions separated; DFDs, schemas and the UI drawn before "
             "code."),
            ("Development",
             "Python 3.12 services + React 19 dashboard; four local models; "
             "**56,789 lines across 182 files**, of which 19,850 are tests."),
            ("Testing",
             "Contract, regression and datastore layers; the whole chain is "
             "**one pass-or-fail command in 73 s**."),
            ("Deployment",
             "The same code in three shapes: host processes, Docker Compose "
             "and Kubernetes with queue-depth autoscaling."),
            ("Evaluation & Audit",
             "Six independent audit passes, four measurement studies, and an "
             "evidence class attached to every claim.")]
    y = TOP + 0.86
    for i, (h1, h2) in enumerate(rows):
        col, row = i % 2, i // 2
        x = L + col * (5.92 + 0.25)
        yy = y + row * 1.00
        sh = card(s, x, yy, 5.92, 0.92, fill=CARD if row % 2 == 0 else CARD2)
        rect(s, x, yy + 0.05, 0.07, 0.82, fill=BLUE if col == 0 else TEAL_L)
        tf = frame_of(sh, 0.22, 0.07)
        para(tf, "%d. %s" % (i + 1, h1), size=12, color=INK, bold=True,
             after=2, first=True, line=0.95)
        para(tf, h2, size=10.5, color=GREY, after=0, line=0.95)
    banner(s, L, 5.32, W, 0.98,
           "Stack — **PostgreSQL + pgvector** (canonical rows and vectors), "
           "**ClickHouse** (analytics events), **Redis Streams** (queues, "
           "cache, progress), **object storage** (raw payloads and reports); "
           "**Ollama / vLLM or a cloud API**, switchable at runtime.",
           fill=INK, size=11.5)

    # ------------------------------------------- 10 architecture
    s = d.body("System Architecture", tag="Chapter 3 · § 3.1")
    shot(s, "fig-dfd0.png", 6.62, TOP + 0.06, w=6.09,
         caption="Context diagram — one read-only pull in, one JSON document out")
    titled_card(s, L, TOP + 0.06, 5.82, 1.62, "A read-only consumer",
                ["One pull of the whole payload — post, comments, "
                 "engagement, reactions, share sample.",
                 "Records are keyed by the **upstream identifier**, so a "
                 "repeated pull is an upsert, not a duplication.",
                 "**Nothing is ever written back**; the platform's own "
                 "coarse score is kept only as a displayed baseline."],
                accent=BLUE, bsize=11)
    titled_card(s, L, TOP + 1.80, 5.82, 1.62, "Its own storage tier",
                ["**PostgreSQL + pgvector** — canonical posts, comments, "
                 "embeddings, watchlist versions.",
                 "**ClickHouse** — analytics events and per-comment "
                 "sentiment at query speed.",
                 "**Redis** — five stream queues, the response cache, usage "
                 "counters and progress events.",
                 "**Object storage** — raw payloads and generated reports."],
                accent=TEAL_L, bsize=11)
    titled_card(s, L, TOP + 3.54, 5.82, 1.20, "Two surfaces on top",
                ["A **51-path HTTP API** generated from the handlers, with "
                 "authentication marked per route.",
                 "An **eleven-tab operator dashboard**, plus nine agents "
                 "reachable over MCP."],
                accent=AMBER, bsize=11)

    # ------------------------------------------- 11 five-stage pipeline
    s = d.body("The five-stage pipeline", tag="Chapter 3 · § 3.2.1")
    steps(s, L, TOP + 0.02, W, 0.60,
          ["1 · Ingestion", "2 · Stage-1 NLP", "3 · Router",
           "4 · Stage-2 LLM", "5 · Assembler"], fill=INK)
    data = [["#", "Stage", "Reads stream", "Principal responsibility"],
            ["1", "Ingestion", "ingestion:queue",
             "Validate against the input schema, normalise Unicode and "
             "script tags, hash and deduplicate, create the job row, enqueue "
             "one message per post"],
            ["2", "Stage-1 NLP", "nlp:stage1:queue",
             "Cheap analysis of the caption **and of every comment**; "
             "sentence embedding; watchlist string matching; sentiment "
             "fusion; a first summary"],
            ["3", "Router", "router:queue",
             "Evaluate the six gates, set the post-level task flags, and "
             "**separately** select the comment set every Stage-2 voter will "
             "read"],
            ["4", "Stage-2 LLM", "llm:stage2:queue",
             "Two concurrent lanes: the **gated** post-level lane and the "
             "**ungated** comment ensemble"],
            ["5", "Assembler", "assembler:queue",
             "Merge both stages, validate against the output schema, fan out "
             "to three stores in parallel, reconcile job counters"]]
    table(s, L, TOP + 0.80, W, data, widths=[0.42, 1.55, 2.05, 8.07],
          font=10.5, hfont=10.5, row_h=0.52, head_h=0.32,
          aligns=["c", None, None, None], first_col_bold=True)
    y = 5.08
    tfh = textbox(s, L, y - 0.02, W, 0.28)
    para(tfh, "Four behaviours implemented once in a shared worker base, so "
              "they hold at every hop", size=12.5, color=INK, bold=True,
         after=0, first=True)
    props = [("Consumer groups on start",
              "Replicas **split** the stream rather than duplicating it."),
             ("Cooperative cancellation",
              "Stages ahead of the assembler drop work on a Redis flag — a "
              "stop costs **one post per stage**."),
             ("Bounded retry, then dead-letter",
              "The attempt counter rides **inside the payload**: a "
              "re-enqueued message would reset a per-id counter forever."),
             ("Replayable progress events",
              "Pub/sub has no backlog, so events are also appended to a "
              "capped list and replayed with a **monotonic sequence**.")]
    cwid = (W - 3 * 0.18) / 4
    for i, (h1, h2) in enumerate(props):
        x = L + i * (cwid + 0.18)
        sh = card(s, x, y + 0.30, cwid, 0.96, fill=CARD2)
        tf = frame_of(sh, 0.13, 0.09)
        para(tf, h1, size=11, color=INK, bold=True, after=2, first=True,
             line=0.95)
        para(tf, h2, size=9.5, color=GREY, after=0, line=0.95)

    # ------------------------------------------------- 12 the router
    s = d.body("The router: two independent decisions",
               tag="Chapter 3 · § 3.2.2")
    banner(s, L, TOP, W, 0.56,
           "Decision A — **does this post earn post-level LLM work?**   "
           "(six gates, any one sufficient)          Decision B — "
           "**which comments will the Stage-2 labellers read?**   "
           "(eligibility → dedup → top-N)", fill=INK, size=12.5,
           align=PP_ALIGN.CENTER)
    data = [["#", "Gate", "Fires when", "Threshold knob (default)"],
            ["1", "Low or absent confidence",
             "Stage-1 overall confidence is below the threshold — **or is "
             "absent entirely**", "ROUTER_CONFIDENCE_THRESHOLD (0.8)"],
            ["2", "Post type unknown or weak",
             "The semantic post type is unset, or its confidence is below "
             "the threshold", "ROUTER_POST_TYPE_CONFIDENCE_THRESHOLD (0.8)"],
            ["3", "Caller asked for a Stage-2 summary",
             "The request explicitly asks for the better summary model",
             "ROUTER_SUMMARY_ROUTES (**false**)"],
            ["4", "Image post with no image verdict",
             "The post carries photographs but the vision stage produced no "
             "sentiment", "not applicable"],
            ["5", "High toxicity",
             "The Stage-1 toxicity score exceeds the threshold",
             "ROUTER_TOXICITY_THRESHOLD (0.7)"],
            ["6", "Long code-mixed text",
             "The caption exceeds the character limit and the script is "
             "Banglish or mixed", "ROUTER_LONG_TEXT_CHARS (1000)"]]
    table(s, L, TOP + 0.72, 7.62, data, widths=[0.34, 1.95, 2.63, 2.70],
          font=9.5, hfont=10, row_h=0.50, head_h=0.32,
          aligns=["c", None, None, None], first_col_bold=True)
    titled_card(s, 8.42, TOP + 0.72, 4.29, 1.62,
                "Decision B — comment selection",
                ["**Eligibility** drops kinds with no text to read: "
                 "emoji-only, link-only, below a word-token floor.",
                 "**Deduplication** sorts by reactions, ties by original "
                 "index, collapses identical normalised texts.",
                 "**Top-N** keeps the survivors; the cap defaults to 0 — "
                 "every unique comment with text."],
                accent=TEAL_L, bsize=10.5, hsize=12)
    titled_card(s, 8.42, TOP + 2.50, 4.29, 1.32,
                "Made once, read by all eight voters",
                ["An earlier design capped **inside** the LLM stance pass "
                 "alone — it still ran seven classifiers over the whole "
                 "thread and left the LLM column empty on exactly the "
                 "comments every cheap head had voted on."],
                accent=AMBER, bsize=10.5, hsize=12)
    banner(s, L, 5.34, W, 0.96,
           "**An absent confidence routes.** In an earlier revision the gate "
           "read a field name Stage 1 never emitted: the guard "
           "short-circuited on every post, a confidence of exactly zero "
           "passed a threshold of 0.65, and the system routed **100 %** of "
           "posts while every document claimed a single-digit rate. The "
           "repair was to make every gate — and the diagnostic log — read "
           "through **named reader functions**, so a rename now fails a "
           "test.        **Gate 3 is off by default**, because any single "
           "gate firing is sufficient: enabled, it made the router a "
           "pass-through and pinned the large-model share at exactly one.",
           fill=INK_D, size=11.5, label="TWO LESSONS LEARNED THE HARD WAY")

    # --------------------------------------------- 13 comment ensemble
    s = d.body("The eight-labeller comment ensemble",
               tag="Chapter 3 · § 3.2.3")
    heads = ["XLM-R\nmultilingual", "mBERT\nsentiment", "BanglaBERT\n(SentNoB)",
             "BanglaBERT\nfine-tune", "ModernBERT", "Twitter-XLM\nsentiment",
             "Emoji &\nlexicon rule"]
    cwid = (8.62 - 6 * 0.10) / 7
    for i, hn in enumerate(heads):
        sh = card(s, L + i * (cwid + 0.10), TOP + 0.30, cwid, 0.78,
                  fill=CARD2, line=BORDER)
        tf = frame_of(sh, 0.05, 0.04)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, hn, size=9.5, color=INK, bold=True, align=PP_ALIGN.CENTER,
             after=0, first=True, line=0.9)
    tfh = textbox(s, L, TOP + 0.02, 8.62, 0.28)
    para(tfh, "Seven small heads — batched on CPU, no per-token cost",
         size=11.5, color=BLUE, bold=True, after=0, first=True)
    sh = card(s, 9.42, TOP + 0.02, 3.29, 1.06, fill=INK, line=INK)
    tf = frame_of(sh, 0.16, 0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "One context-aware LLM stance pass", size=11.5, color=WHITE,
         bold=True, after=2, first=True, line=0.95)
    para(tf, "The only labeller that sees the parent post — and therefore "
             "the only one judging stance *toward* it rather than the "
             "comment's own tone.", size=9.5,
         color=RGBColor(0xD6, 0xE4, 0xF2), after=0, line=0.95)
    banner(s, L, TOP + 1.24, W, 0.66,
           "**Abstention protocol —** a labeller that cannot answer is "
           "**absent** from the vote, not present in it with a fabricated "
           "neutral. Every comment carries the number of voters that "
           "actually spoke; a comment no model read reports as "
           "**uncertain at zero voters**, never as neutral.",
           fill=AMBER, size=12, color=WHITE)
    eq = card(s, L, TOP + 2.06, 5.85, 1.52, fill=CARD)
    tf = frame_of(eq, 0.22, 0.12)
    para(tf, "Combined polarity over the labellers that spoke", size=12,
         color=INK, bold=True, after=5, first=True)
    para(tf, "s(c)  =  ( 1 / nᶜ )  ·  Σ  wᵢ · vᵢ(c)", size=16,
         color=INK, bold=True, align=PP_ALIGN.CENTER, after=5,
         font="Cambria Math")
    para(tf, "The sum runs over the nᶜ labellers in V(c) that returned a "
             "verdict; vᵢ(c) ∈ {−1, 0, +1}, with the weights "
             "renormalised so s(c) stays in [−1, +1] however many "
             "labellers abstained.", size=10, color=GREY, after=0, line=0.95)
    titled_card(s, 6.86, TOP + 2.06, 5.85, 1.52,
                "What the operator sees, per comment",
                ["The fused label **and** the eight per-source verdicts side "
                 "by side, the agreement fraction, the voter count, the "
                 "**method** that actually ran and the comment **kind** "
                 "(emoji · short · substantive)."],
                accent=TEAL_L, bsize=11, hsize=12)
    banner(s, L, 5.10, W, 1.20,
           "Why eight rather than one — a fused single label makes a member "
           "model that failed to load **unfalsifiable**: it is silently "
           "replaced by a default and the reported agreement stops meaning "
           "anything. Showing the roster is what makes the number "
           "checkable [50].", fill=INK_D, size=11.5)

    # ---------------------------------------- 14 stance + agentic layer
    s = d.body("Watchlist target stance, and the agentic insight layer",
               tag="Chapter 3 · § 3.2.4 – 3.2.6")
    titled_card(s, L, TOP, 5.85, 2.28,
                "Watchlist-driven target stance",
                ["An operator-supplied, **versioned** entity list; the "
                 "matcher is built for the three-script spelling problem "
                 "— Bangla script, several inconsistent romanizations, and "
                 "English.",
                 "Per-entity verdicts land in a field kept **separate from "
                 "document sentiment**, and are never summed or charted "
                 "with it.",
                 "**Zero additional LLM calls** — matched entities ride "
                 "inside a prompt that is already being sent.",
                 "A target matching **zero** comments is logged as the "
                 "probable coverage bug it almost certainly is."],
                accent=TEAL, bsize=11)
    titled_card(s, 6.86, TOP, 5.85, 2.28,
                "Nine agents over eighteen typed MCP tools",
                ["Three Model Context Protocol servers [38] expose the "
                 "system's **own structured analytics** — not raw text — as "
                 "typed tools.",
                 "Agents run a ReAct-style loop [36] under **per-run tool "
                 "and token budgets**, outside the per-post path.",
                 "Briefing, stance, theme, anomaly and comparison agents all "
                 "read the same persisted analysis the dashboard shows."],
                accent=BLUE, bsize=11)
    tfh = textbox(s, L, TOP + 2.44, W, 0.28)
    para(tfh, "The runner is hardened against failure modes observed in live "
              "runs — each one now pinned by a regression test", size=12.5,
         color=INK, bold=True, after=0, first=True)
    guards = [("Untrusted-data delimiters",
               "Tool results are wrapped, with **forged delimiters "
               "neutralised** before the model sees them [43]."),
              ("Citation verification",
               "Every citation is checked against **what the model was "
               "actually shown**, not against the corpus."),
              ("Non-answer rejection",
               "Four distinct shapes of answer-that-answers-nothing are "
               "detected and refused."),
              ("Repeat-call and budget guards",
               "Byte-identical repeat tool calls are refused; every run is "
               "bounded in tools and tokens.")]
    cwid = (W - 3 * 0.18) / 4
    for i, (h1, h2) in enumerate(guards):
        x = L + i * (cwid + 0.18)
        sh = card(s, x, TOP + 2.74, cwid, 1.00, fill=CARD2)
        tf = frame_of(sh, 0.13, 0.09)
        para(tf, h1, size=11, color=INK, bold=True, after=2, first=True,
             line=0.95)
        para(tf, h2, size=9.5, color=GREY, after=0, line=0.95)
    banner(s, L, 5.24, W, 1.06,
           "A briefing that quoted invented tables was caught by this layer "
           "and is reported in the walkthrough: the agent had retrieved "
           "**nothing**, and the run record proved it. Groundedness and "
           "citation accuracy are **hardened and tested, but not scored** — "
           "they sit at *works, unmeasured* in the evidence ledger, and the "
           "scoring programme is future work F7.",
           fill=INK_D, size=11.5, label="WHAT THE HARDENING IS, AND IS NOT")

    # ------------------------------------------ 15 verification result
    s = d.body("Results — verification of the whole chain",
               tag="Chapter 4 · § 4.2.2")
    for i, (v, lab, col) in enumerate([
            ("1,420", "Python tests", INK),
            ("96", "dashboard unit tests", TEAL),
            ("138", "security tests · 8 suites", INK),
            ("73 s", "whole chain, one command", TEAL),
            ("51", "documented API paths", INK)]):
        kpi(s, L + i * 2.45, TOP + 0.02, 2.26, 0.94, v, lab, fill=col,
            vsize=24)
    tfh = textbox(s, L, TOP + 1.16, W, 0.28)
    para(tfh, "Three layers, because the failures that dominated this "
              "project are not failures a person notices by looking",
         size=12.5, color=INK, bold=True, after=0, first=True)
    layers = [("Contract tests",
               "Assert that the JSON crossing every boundary validates "
               "against its schema — and several **grep the source** to "
               "assert structural invariants: that every worker checks the "
               "cancellation flag, that the dashboard reads only fields the "
               "API returns, that the autoscaler manifests name exactly the "
               "streams the workers create. **They fail on a rename, which "
               "is their purpose.**", BLUE),
              ("Regression tests",
               "Pin every defect found in the six audit passes, so a "
               "repaired defect that returns **fails the build**. Eight "
               "distinct agent-runner failure modes and four "
               "silent-degradation defects each have one.", TEAL),
              ("Datastore tests",
               "Run against **throwaway PostgreSQL and Redis containers** "
               "started per session, so a default run can never damage a "
               "developer's environment; two heavier suites sit behind "
               "explicit markers.", AMBER)]
    cwid = (W - 2 * 0.22) / 3
    for i, (h1, h2, acc) in enumerate(layers):
        titled_card(s, L + i * (cwid + 0.22), TOP + 1.46, cwid, 1.90, h1,
                    [h2], accent=acc, bsize=11)
    banner(s, L, 4.86, W, 0.72,
           "Access control is tested **separately from function**, because "
           "the two fail differently: token issue and verification, the "
           "hashed API-key store, single-use streaming tickets, cross-tenant "
           "reads, the fail-closed residency policy and prompt-injection "
           "containment — against both the permitted **and** the refused "
           "case in each.", fill=INK, size=11.5)
    tfn = textbox(s, L, 5.72, W, 0.6)
    para(tfn, "A green suite is not a quality claim — and this report does "
              "not use it as one. These are **correctness and contract** "
              "tests; what they buy is that a claim in this report which "
              "stopped being true would break a test.", size=12,
         color=AMBER, bold=True, after=0, first=True, line=1.0)

    # ------------------------------------------- 16 routing behaviour
    s = d.body("Results — routing behaviour", tag="Chapter 4 · § 4.2.4")
    tfh = textbox(s, L, TOP, 5.85, 0.3)
    para(tfh, "Posts routed for post-level LLM work, by Stage-1 engine",
         size=12, color=INK, bold=True, after=0, first=True)
    bar_chart(s, L - 0.14, TOP + 0.28, 6.0, 2.36,
              ["Keyword-stub\nStage 1", "LLM Stage 1\n(as shipped)"],
              [("Percent of posts routed", (74, 16))], [BLUE, ],
              num_fmt='0"%"', label_size=13, gap=150, y_max=100)
    data = [["Quantity", "Keyword stub", "As shipped"],
            ["Posts routed for post-level work", "32 of 43  (74 %)",
             "7 of 43  (16 %)"],
            ["Post-level calls", "124  (15.3 %)", "22  (4.2 %)"],
            ["Comment-level calls", "689  (84.7 %)", "507  (95.8 %)"],
            ["   · Stage-1 labelling — all 43 posts, the gate cannot reduce it",
             "368", "368"],
            ["   · Stage-2 comment stance — routed posts only", "321", "139"],
            ["Total calls per corpus run", "813", "529"],
            ["**Share of all calls the gate governs**", "**54.7 %**",
             "**30.4 %**"]]
    table(s, 6.86, TOP + 0.28, 5.85, data, widths=[3.55, 1.15, 1.15],
          font=9.5, hfont=10, row_h=0.335, head_h=0.32,
          aligns=[None, "c", "c"], bold_rows=(7,))
    banner(s, L, 4.14, W, 0.86,
           "Identical router code, identical corpus, identical thresholds — "
           "**only the Stage-1 engine differs.** The keyword stub confidently "
           "types 27 of 43 posts and escalates the remaining 74 %; the LLM "
           "Stage 1 types 42 of 43 and escalates 16 %. So the routing rate "
           "measures **Stage-1 quality, not cost efficiency** — a rising "
           "rate should be alerted on as a quality regression.",
           fill=INK, size=11.5, label="READING THE CHART")
    titled_card(s, L, 5.14, 5.85, 1.16, "The gate selects the right posts",
                ["The 7 routed posts hold **3,404 of the 8,713** comments "
                 "that reach a model — the gate escalates the large, "
                 "contentious threads, which a routing rate quoted on its "
                 "own cannot show."], accent=GREEN, bsize=11, hsize=12)
    titled_card(s, 6.86, 5.14, 5.85, 1.16,
                "One earlier estimate was off by 6×",
                ["Filtering emoji-only comments was expected to save ~17 % "
                 "of the comment bill. They are **2.8 %** of the corpus; the "
                 "17 % was the share taking the *fast path*. Still correct "
                 "to filter — but not a cost lever."],
                accent=AMBER, bsize=11, hsize=12)

    # --------------------------------------------- 17 the cost split
    s = d.body("Results — where the large-model calls actually go",
               tag="Chapter 4 · § 4.2.4")
    bar_chart(s, L - 0.14, TOP + 0.02, 7.2, 2.72,
              ["Post-level", "Stage-1 comment\nlabelling",
               "Stage-2 comment\nstance"],
              [("Keyword-stub Stage 1", (124, 368, 321)),
               ("LLM Stage 1 (shipped)", (22, 368, 139))],
              [BLUE_L, INK], num_fmt='0', legend=True, label_size=11,
              y_max=420)
    eq = card(s, 7.90, TOP + 0.02, 4.81, 1.42, fill=CARD)
    tf = frame_of(eq, 0.20, 0.11)
    para(tf, "Share of spend the gate governs", size=12, color=INK,
         bold=True, after=4, first=True)
    para(tf, "γ  =  C_gated  /  ( C_gated + C_comment )", size=13,
         color=INK, bold=True, align=PP_ALIGN.CENTER, after=5,
         font="Cambria Math")
    para(tf, "C_gated counts the calls the gate releases; C_comment counts "
             "the per-comment calls, which run for every post whether the "
             "gate routed it or not.", size=10, color=GREY, after=3,
         line=0.95)
    para(tf, "Measured:  **γ = 0.304** as shipped,  **γ = 0.547** under the "
             "keyword baseline.", size=10.5, color=INK, after=0, line=0.95)
    titled_card(s, 7.90, TOP + 1.58, 4.81, 1.16,
                "The middle pair of bars is the finding",
                ["Stage-1 comment labelling runs for **every** post whether "
                 "the gate routed it or not, so it is identical in both "
                 "configurations — the routing gate **cannot reduce it at "
                 "all**."], accent=AMBER, bsize=11, hsize=12)
    banner(s, L, 4.02, W, 1.06,
           "It was wrong **in fact** — two rules read field names Stage 1 "
           "never emitted, one read a placeholder Stage 2 fills, and one "
           "compared a language value against a marker stored elsewhere: net, "
           "one rule always fired and three could never fire. It was wrong "
           "**in framing** — even a correctly functioning gate governs a "
           "minority of spend once every comment is labelled.",
           fill=RED, size=11.5, color=WHITE,
           label="THE PROJECT'S OWN HEADLINE CLAIM WAS WRONG, TWICE OVER",
           )
    titled_card(s, L, 5.22, 5.85, 1.08, "The claim that survives",
                ["**Cheap NLP filters which comments and which posts deserve "
                 "a large model.** Both halves matter — and the second is "
                 "now the bigger one."], accent=GREEN, bsize=11.5, hsize=12)
    titled_card(s, 6.86, 5.22, 5.85, 1.08,
                "The transferable lesson",
                ["**When the routed unit contains many sub-items that each "
                 "need work, routing the container is not the cost lever.** "
                 "The published cascades route single queries, so they never "
                 "meet this."], accent=INK, bsize=11.5, hsize=12)

    # ------------------------------------------------- 18 retrieval
    s = d.body("Results — retrieval evaluation", tag="Chapter 4 · § 4.2.5")
    tfh = textbox(s, L, TOP, 5.9, 0.28)
    para(tfh, "recall@10 before and after real embeddings", size=12,
         color=INK, bold=True, after=0, first=True)
    bar_chart(s, L - 0.14, TOP + 0.26, 6.05, 2.30,
              ["vector", "lexical", "hybrid"],
              [("Hash-stub vectors", (0.1875, 0.3750, 0.5312)),
               ("Real embeddings", (0.6250, 0.7500, 0.8750))],
              [BLUE_L, INK], num_fmt='0.000', legend=True, label_size=9.5,
              y_max=1.0)
    tfh = textbox(s, 6.86, TOP, 5.85, 0.28)
    para(tfh, "recall@10 stratified by query-to-target lexical overlap",
         size=12, color=INK, bold=True, after=0, first=True)
    bar_chart(s, 6.72, TOP + 0.26, 6.05, 2.30,
              ["overlap ≥ 0.5\n(n = 12)", "overlap < 0.5\n(n = 20)",
               "overlap = 0\n(n = 8)"],
              [("lexical", (1.000, 0.600, 0.000)),
               ("vector (dense)", (0.583, 0.650, 0.500)),
               ("chunked (fused)", (1.000, 0.800, 0.500))],
              [BLUE_L, TEAL_L, INK], num_fmt='0.000', legend=True,
              label_size=9, y_max=1.15)
    data = [["Configuration", "recall@10 (hash)", "recall@10 (real)",
             "MRR@10", "nDCG@10"],
            ["vector (dense only)", "0.1875", "0.6250", "0.2869", "0.3667"],
            ["lexical (full text fused with trigram)", "0.3750", "0.7500",
             "0.6503", "0.6748"],
            ["chunk (dense over passage chunks)", "not measured", "0.6562",
             "0.3014", "0.3846"],
            ["hybrid (dense fused with lexical)", "0.5312", "0.8750",
             "0.6108", "0.6759"],
            ["chunked (chunk fused with full text and trigram)", "not measured",
             "**0.8750**", "**0.6482**", "**0.7041**"]]
    table(s, L, 4.02, 7.34, data, widths=[3.5, 1.1, 1.1, 0.9, 0.9],
          font=9.5, hfont=9.5, row_h=0.315, head_h=0.42,
          aligns=[None, "c", "c", "c", "c"], bold_rows=(5,))
    titled_card(s, 8.16, 4.02, 4.55, 2.02, "Three findings",
                ["Hash vectors were **indistinguishable from chance** — ten "
                 "random posts of fifty score 0.20 by construction.",
                 "**Chunking improves ranking, not recall** here: only 14 of "
                 "50 posts exceed 600 characters.",
                 "The lexical arm's jump was a **bug fix**: the query builder "
                 "conjoined every term, so the full-text half matched "
                 "**zero** rows — covered by a passing test."],
                accent=AMBER, bsize=10.5, hsize=12)
    tfn = textbox(s, L, 6.06, 7.34, 0.3)
    para(tfn, "32 known-item queries, one marked relevant post each, k = 10, "
              "50-post corpus.   Lower bounds: every other on-topic post "
              "counts as a miss, and the decisive bucket holds 8 queries — "
              "**the direction is consistent; the set is too small to call "
              "it.**", size=9.5, color=MUTED, after=0, first=True, line=0.95)

    # ------------------------------------------------ 19 provenance
    s = d.body("Results — label provenance and ensemble coverage",
               tag="Chapter 4 · § 4.2.6")
    pie_chart(s, L - 0.10, TOP, 5.3, 2.9,
              ["Model or LLM output  (2,953)",
               "Deterministic stub or emoji-lexicon rule  (7,319)"],
              (2953, 7319), [INK, BLUE_L], num_fmt='0.0%')
    titled_card(s, 5.60, TOP + 0.04, 7.11, 1.46,
                "What the measurement found",
                ["Measured over all **10,272** comments in the repository's "
                 "**default** configuration. The larger slice is partly a "
                 "fourteen-word emoji-and-lexicon rule and mostly a "
                 "deterministic function of a hash of the first fifty "
                 "characters. It is reproducible, and **it is not "
                 "sentiment**.",
                 "Before provenance was recorded, the same run reported "
                 "**8,513 model inferences in a configuration where zero "
                 "models had been loaded.**"],
                accent=AMBER, bsize=11, hsize=12)
    tfh = textbox(s, 5.60, TOP + 1.66, 7.11, 0.28)
    para(tfh, "Four reporting mechanisms — now permanent parts of the output "
              "contract", size=12, color=INK, bold=True, after=0, first=True)
    mechs = [("method", "names the engine that actually ran, per comment"),
             ("kind", "distinguishes emoji · short · substantive text"),
             ("provenance block", "per post: what share of labels were "
              "model inferences"),
             ("voters that spoke", "beside every label and agreement figure")]
    for i, (h1, h2) in enumerate(mechs):
        col, row = i % 2, i // 2
        x = 5.60 + col * (3.50 + 0.11)
        yy = TOP + 1.96 + row * 0.60
        sh = card(s, x, yy, 3.50, 0.52, fill=CARD2)
        tf = frame_of(sh, 0.12, 0.04)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, h1, size=10.5, color=INK, bold=True, after=1, first=True,
             line=0.9)
        para(tf, h2, size=9.5, color=GREY, after=0, line=0.9)
    banner(s, L, 4.46, W, 0.62,
           "The remedy was **not to hide the slice but to label it** — which "
           "makes the accurate statement available: 28.7 % of labels are "
           "model output, with the remainder itemised. With the full "
           "ensemble roster provisioned, that figure is 100 %.",
           fill=INK, size=12)
    data = [["Quantity", "All-cheap baseline", "Shipped hybrid",
             "LLM-everywhere"],
            ["Posts routed for post-level LLM work", "32 of 43  (74 %)",
             "7 of 43  (16 %)", "43 of 43, by definition"],
            ["Total LLM calls per corpus run", "813", "529",
             "not run — bounded below by 1/post + 1/comment batch"],
            ["Independent verdicts per analysed comment", "1",
             "up to **8**, with abstention", "1"],
            ["Can the output state what produced each label?", "no",
             "**yes**", "trivially — one producer"]]
    table(s, L, 5.16, W, data, widths=[3.6, 2.1, 2.4, 3.99],
          font=9.5, hfont=10, row_h=0.25, head_h=0.27,
          aligns=[None, "c", "c", "c"])

    # -------------------------------------------------- 20 novelty
    s = d.body("Novelty of the Work", tag="Contribution")
    nov = [("A cascade that reports what its gate is worth",
            "The router makes **two separate decisions**, and the system "
            "publishes γ — the share of spend the gate governs — from its "
            "own counters. No published cascade routes a *container* of "
            "sub-items, so none reports this quantity at all.", INK),
           ("Abstention instead of a fabricated default",
            "Eight independent verdicts per comment with the roster shown "
            "side by side; a comment no model read is **uncertain at zero "
            "voters**, never neutral. It makes the reported agreement "
            "falsifiable.", TEAL),
           ("Alias-aware target stance for three scripts",
            "Operator-supplied, versioned watchlist; per-entity verdicts at "
            "**zero additional model calls**, in a field kept separate from "
            "sentiment — and a zero-match target is logged as the coverage "
            "bug it probably is.", AMBER),
           ("Agents over the system's own analytics, hardened",
            "Nine agents, eighteen typed MCP tools, three servers — with "
            "untrusted-data delimiters, forged-delimiter neutralisation, "
            "citation verification against what the model was shown, and "
            "non-answer rejection.", BLUE),
           ("Provenance on every value, evidence class on every claim",
            "Twenty-eight capabilities each recorded once as *measured*, "
            "*works-but-unmeasured*, *unexercised* or *out of scope* — so no "
            "claim rests on a stronger basis than the ledger records.",
            GREEN),
           ("A disconfirmation reported as a result",
            "The project's own headline claim did not survive measurement, "
            "and the report says so. The generalisable finding — **silent "
            "successes, not wrong answers, are what break systems of this "
            "shape** — came out of that.", RED)]
    cwid, chh = 3.87, 1.68
    for i, (h1, h2, acc) in enumerate(nov):
        col, row = i % 3, i // 3
        titled_card(s, L + col * (cwid + 0.24), TOP + 0.02 + row * (chh + 0.20),
                    cwid, chh, h1, [h2], accent=acc, bsize=10.5, hsize=12)
    banner(s, L, 5.06, W, 1.24,
           "Twenty-one of twenty-three capability rows in the feature "
           "comparison are **Yes** where the four comparable classes of "
           "system are **No**, *Partial* or *Not applicable*. The two rows "
           "that are not — the image modality and labelling accuracy — are "
           "the ones that qualify the rest, and they are carried at their "
           "evidence class throughout the report rather than quietly "
           "dropped.", fill=INK_D, size=11.5)

    # ------------------------------------------------ 21 sample dataset
    s = d.body("Sample Dataset — the working corpus",
               tag="Chapter 4 · § 4.2.1")
    data = [["Property", "Value"],
            ["Posts", "50, across 25 monitoring campaigns"],
            ["Media type", "37 photo-with-text, 7 photo-only, 6 text-only"],
            ["Posts with a caption", "43 of 50 — the other 7 have a null "
             "caption and no reachable image"],
            ["Caption length", "median 245 characters, maximum 5,427; "
             "14 of 50 exceed 600"],
            ["Publication window", "29 April to 19 May 2026"],
            ["Comments shipped in the payload", "**10,272**"],
            ["Comments the platform reports", "274,126  →  corpus coverage "
             "**3.75 %**"],
            ["Stored comments per post", "minimum 87, median 98, maximum "
             "**2,857**"],
            ["Comment reactions", "mean 3.97, maximum 946; 64.1 % have zero"],
            ["Emoji-only comments (pipeline's own kind classifier)", "2.8 %"],
            ["Threading", "0 comments carry a parent id, yet 1,111 report a "
             "non-zero reply count"],
            ["Post reactions · shares", "794,373 · 93,903"],
            ["Image references", "69 across 44 posts — all relative "
             "object-storage keys, **objects absent**"]]
    table(s, L, TOP + 0.02, 7.18, data, widths=[3.05, 4.13],
          font=9.5, hfont=10, row_h=0.302, head_h=0.30, aligns=[None, None])
    tfh = textbox(s, 8.00, TOP + 0.02, 4.71, 0.28)
    para(tfh, "Script composition of the 10,272 comments", size=12,
         color=INK, bold=True, after=0, first=True)
    pie_chart(s, 8.00, TOP + 0.26, 4.71, 2.60,
              ["Bangla script only — 81.4 %", "Latin script only — 11.1 %",
               "Both scripts, code-mixed — 4.6 %",
               "Neither, emoji or symbols — 2.9 %"],
              (8357, 1143, 473, 299), [INK, BLUE, TEAL_L, BLUE_L],
              labels=False)
    titled_card(s, 8.00, TOP + 3.00, 4.71, 2.06,
                "Three consequences, stated up front",
                ["**The comment sample is not random.** It is the "
                 "platform's engagement-ordered set, so what the data "
                 "supports is a claim about *the most-engaged N comments per "
                 "post* — and that is the phrasing the report uses.",
                 "**Post sentiment is a text measurement.** No image bytes "
                 "are reachable, so the vision term has never carried a "
                 "value; fusion renormalises over the terms that did.",
                 "**The 7 caption-less posts stay in.** They carry 1,307 "
                 "comments; excluding them measured a population the running "
                 "system never processes."],
                accent=AMBER, bsize=10, hsize=12)
    tfn = textbox(s, L, 5.68, 7.18, 0.7)
    para(tfn, "Five posts exceed 100 % coverage — up to 112 stored comments "
              "against a reported 42, which cannot be replies inflating the "
              "numerator because no comment carries a parent id. The "
              "pipeline **clamps coverage to 1.0 and raises an explicit "
              "anomaly flag** rather than rendering an impossible figure.",
         size=10.5, color=GREY, after=0, first=True, line=1.0)

    # ---------------------------------------------- 22 expected output
    s = d.body("Expected Output — one document per thread",
               tag="Chapter 4 · § 4.3.1")
    shot(s, "fig-shot-canonical.png", L, TOP + 0.02, w=7.18,
         caption="The canonical result document for one post, as the API "
                 "returns it")
    titled_card(s, 7.96, TOP + 0.02, 4.75, 2.30, "What the document carries",
                ["**Post block** — fused sentiment with its confidence, "
                 "semantic post type, topics, entities, summary and insight.",
                 "**Comment array** — per comment: fused label, agreement, "
                 "**voters that spoke**, the eight per-source verdicts, "
                 "method and kind.",
                 "**Target stance** — per watchlist entity, in its own "
                 "field, never summed with sentiment.",
                 "**Router decision** — the branch taken and every reason "
                 "string that fired.",
                 "**Provenance and coverage** — inferred share, comment "
                 "coverage, and any anomaly flags raised."],
                accent=BLUE, bsize=10.5, hsize=12)
    titled_card(s, 7.96, TOP + 2.46, 4.75, 1.30, "The contract, enforced",
                ["Validated against the output schema **before** it is "
                 "written, then fanned out to PostgreSQL, ClickHouse and "
                 "object storage in parallel — with the job counters "
                 "reconciled from the rows that actually landed."],
                accent=TEAL_L, bsize=10.5, hsize=12)
    banner(s, 7.96, TOP + 3.92, 4.75, 1.12,
           "The upstream platform's own coarse score is preserved **beside** "
           "the recomputed one, never overwritten — so a consumer can always "
           "see both, and the analysis layer stays auditable against its "
           "source.", fill=INK, size=11)

    # ------------------------------------------------- 23 web interface I
    s = d.body("Web Interface — corpus, comments and the bill",
               tag="Chapter 4 · § 4.3.1")
    tabs = ["Overview", "Posts", "Jobs", "Reports", "Search", "Agents",
            "Chat", "Pipeline", "Trace", "Warnings", "Logs"]
    tw = (W - 10 * 0.09) / 11
    for i, tb_ in enumerate(tabs):
        sh = rect(s, L + i * (tw + 0.09), TOP, tw, 0.34,
                  fill=INK if i in (0, 1) else CARD2,
                  line=BORDER if i not in (0, 1) else INK,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28)
        tf = frame_of(sh, 0.04, 0.01)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, tb_, size=10, color=WHITE if i in (0, 1) else INK,
             bold=True, align=PP_ALIGN.CENTER, after=0, first=True)
    shot(s, "fig-shot-economics.png", 1.11, TOP + 0.50, w=5.35,
         caption="Token economics — spend by lane, by model, and projected")
    shot(s, "fig-shot-ensemble.png", 6.86, TOP + 0.50, w=5.35,
         caption="The eight-labeller ensemble, per comment, side by side")
    banner(s, L, 5.56, W, 0.78,
           "In the run shown, five lanes divide 941 calls — post 10, comment "
           "284, Stage-1 83, interactive 523, agent 41. By tokens the "
           "**comment lane takes 49 %** and the interactive lane 35 %, while "
           "**the post lane — the only one the router's gate governs — takes "
           "under one percent.** That is § 4.2.4's finding arrived at from "
           "the usage ledger instead of the router's counters.",
           fill=INK_D, size=11.5, label="THE SAME FINDING, FROM A SECOND SOURCE")

    # ------------------------------------------------ 24 web interface II
    s = d.body("Web Interface — running it, watching it, asking it",
               tag="Chapter 4 · § 4.3.1")
    shot(s, "fig-shot-pipeline-live2.png", L, TOP, w=3.88,
         caption="Pipeline — live queue depths and per-stage throughput")
    shot(s, "fig-ui-trace.png", 4.72, TOP, w=3.88,
         caption="Trace — one post's five stages as they happen")
    shot(s, "fig-shot-agent-run.png", 8.83, TOP, w=3.88,
         caption="Agents — a run with its tool calls and citations")
    y = TOP + 2.86
    cells = [("One command, eleven processes",
              "The orchestrator syncs dependencies, starts four datastores, "
              "creates the ClickHouse tables, **verifies the encoder "
              "produces real vectors rather than hash stubs**, resets the "
              "consumer groups and launches eleven host processes.", BLUE),
             ("A fabrication, caught in the act",
              "The Chat tab is where a briefing that quoted invented tables "
              "was caught: the agent had retrieved **nothing**, and the run "
              "record in Redis proved it. The guard is now code.", RED),
             ("Two operational gaps, recorded not hidden",
              "The event stream has a **five-minute server-side ceiling**, "
              "after which Trace reports the run as failed while the "
              "pipeline finishes correctly; the monitoring stack ships a "
              "correct scrape and **no provisioned dashboards**.", AMBER)]
    cwid = (W - 2 * 0.24) / 3
    for i, (h1, h2, acc) in enumerate(cells):
        titled_card(s, L + i * (cwid + 0.24), y, cwid, 1.34, h1, [h2],
                    accent=acc, bsize=10.5, hsize=12)
    tfn = textbox(s, L, y + 1.46, W, 0.3)
    para(tfn, "Every capture is the running system on the development tier "
              "with the corpus loaded — **nothing here is a mock-up**; where "
              "a panel shows a number, that number is what the system "
              "computed.", size=11.5, color=INK, italic=True, after=0,
         first=True)

    # ------------------------------------------------------- 25 video
    s = d.body("Demonstration Video", tag="Live system")
    fw, fh = 7.62, 4.29
    frame = card(s, L, TOP + 0.06, fw, fh, fill=RGBColor(0x10, 0x1B, 0x2B),
                 line=INK, radius=0.02)
    rect(s, L, TOP + 0.06, fw, 0.30, fill=INK_D)
    tfb = textbox(s, L + 0.16, TOP + 0.11, fw - 0.32, 0.22)
    para(tfb, "selective-intelligence — end-to-end demonstration", size=9.5,
         color=BLUE_L, after=0, first=True)
    cx, cy, cr = L + fw / 2, TOP + 0.06 + fh / 2, 0.52
    rect(s, cx - cr, cy - cr - 0.18, cr * 2, cr * 2, fill=WHITE, line=WHITE,
         shape=MSO_SHAPE.OVAL)
    tri = rect(s, cx - 0.20, cy - 0.44, 0.44, 0.50, fill=INK, line=INK,
               shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    tri.rotation = 90
    tfc = textbox(s, L + 0.4, cy + 0.52, fw - 0.8, 0.9)
    para(tfc, "Place the demonstration video here", size=15, color=WHITE,
         bold=True, align=PP_ALIGN.CENTER, after=4, first=True)
    para(tfc, "PowerPoint ▸ Insert ▸ Video ▸ This Device… — then size the "
              "video to this frame and delete this placeholder group.",
         size=10.5, color=BLUE_L, align=PP_ALIGN.CENTER, after=0, line=0.95)
    tfh = textbox(s, 8.42, TOP + 0.06, 4.29, 0.3)
    para(tfh, "What the recording walks through", size=13, color=INK,
         bold=True, after=0, first=True)
    walk = ["**One command** brings up four datastores and eleven processes, "
            "with the encoder preflight visible.",
            "**Ingest** the 50-post payload; the job appears with its "
            "counters.",
            "**Pipeline** tab: queue depths move stage by stage as the "
            "thread is analysed.",
            "**Trace** one post: ingestion → Stage-1 → router decision and "
            "its reason strings → Stage-2 → assembler.",
            "**Posts** tab: open a thread and read the eight per-source "
            "verdicts, voters and abstentions on one comment.",
            "**Warnings**: a watchlist target with per-entity stance, and a "
            "zero-match target flagged.",
            "**Search** in both modes, then an **agent run** with its tool "
            "calls and verified citations.",
            "**Economics**: the five lanes, and the post lane at under one "
            "percent of tokens."]
    tfw = textbox(s, 8.42, TOP + 0.40, 4.29, 4.0)
    for i, w_ in enumerate(walk):
        para(tfw, w_, size=10.5, color=GREY, bullet="▸", bullet_color=BLUE,
             after=6, line=0.95, first=(i == 0))
    banner(s, 8.42, 5.10, 4.29, 1.20,
           "Recorded on the development tier of the deployment figure, with "
           "the corpus of § 4.2.1 already loaded — the same configuration "
           "every screenshot in Chapter 4 was taken from.", fill=INK,
           size=10.5, label="RECORDING CONDITIONS")

    # ---------------------------------------------------- 26 conclusion
    s = d.body("Conclusion", tag="Chapter 6 · § 6.1")
    banner(s, L, TOP, W, 0.80,
           "Given a post together with its comment thread, in a setting that "
           "mixes Bangla script, English and romanized Banglish — can a "
           "system decide **per unit of work** how much intelligence that "
           "unit deserves, spend the expensive model only where cheap models "
           "cannot answer, and prove afterwards where the money went and "
           "what produced every label?", fill=INK, size=12.5,
           label="THE QUESTION THIS PROJECT ASKED")
    tfa = textbox(s, L, TOP + 0.92, W, 0.36)
    para(tfa, "The answer is yes — and not in the way the project originally "
              "believed.", size=15, color=AMBER, bold=True, after=0,
         first=True)
    titled_card(s, L, TOP + 1.34, 3.87, 2.06, "What was built",
                ["A five-stage, queue-decoupled microservice, measured "
                 "rather than asserted.",
                 "A six-gate router making **two** decisions — post-level "
                 "spend, and comment selection.",
                 "Eight labellers per comment with explicit abstention.",
                 "Alias-aware watchlist stance at zero extra calls, and nine "
                 "agents over eighteen typed MCP tools."],
                accent=BLUE, bsize=10.5)
    titled_card(s, L + 4.11, TOP + 1.34, 3.87, 2.06, "What was measured",
                ["**16 %** of posts route under the shipped engine, against "
                 "**74 %** for a keyword baseline —",
                 "yet the gate governs only **30 %** of large-model calls, "
                 "because 85–96 % of calls are comment-level and comment "
                 "labelling runs for every post regardless.",
                 "Hybrid retrieval reaches **0.8750** recall@10 where "
                 "hash-seeded vectors score **0.1875**."],
                accent=TEAL, bsize=10.5)
    titled_card(s, L + 8.22, TOP + 1.34, 3.87, 2.06,
                "What transfers to the next system",
                ["The consequential failures here were not wrong answers but "
                 "**silent successes**: a gate reporting plausible behaviour "
                 "with four of six rules inert; a retrieval arm matching zero "
                 "rows under a passing test; a server discarding 11,000 "
                 "prompt tokens and reporting only what it evaluated."],
                accent=AMBER, bsize=10.5)
    banner(s, L, TOP + 3.56, W, 0.78,
           "The engineering response is the deliverable's real content — "
           "**named reader functions, single-source constants, provenance on "
           "every value, abstention instead of a default vote, and an "
           "evidence class on every published claim.**", fill=INK_D,
           size=12.5)
    for i, (v, lab) in enumerate([("56,789", "lines of Python · 182 files"),
                                  ("19,850", "of them tests"),
                                  ("1,420 + 96", "Python + dashboard tests"),
                                  ("73 s", "whole verification chain")]):
        kpi(s, L + i * 3.055, 5.48, 2.86, 0.82, v, lab,
            fill=TEAL if i % 2 else INK, vsize=20, lsize=10)

    # -------------------------------------- 27 limitations & future work
    s = d.body("Limitations, and the work that closes them",
               tag="Chapter 6 · § 6.2 – 6.3")
    tfh = textbox(s, L, TOP, 5.92, 0.3)
    para(tfh, "Six limitations — five constrain the evidence, the sixth is a "
              "scope decision", size=12.5, color=INK, bold=True, after=0,
         first=True)
    lims = [("L1", "The corpus is small and its comment sample is biased — "
             "50 posts, one platform, a three-week window; 10,272 comments "
             "are an engagement-ordered **3.75 %** slice."),
            ("L2", "The multimodal claim is **unexercised**: 44 of 50 posts "
             "carry images, but the 69 object keys have no objects."),
            ("L3", "Retrieval is measured against a **machine-derived** "
             "query set — no human relevance judgement, one relevant post "
             "per query."),
            ("L4", "Scale is code-complete and **unbenchmarked**: no "
             "multi-replica run has been timed in a cluster."),
            ("L5", "The agent layer's answers are **unmeasured** — the "
             "runner is hardened and tested; groundedness is not scored."),
            ("L6", "**Labelling accuracy is out of scope.** The sampler and "
             "scorer ship; the gold file holds **zero adjudicated labels**, "
             "deliberately — labels seeded from a model in this repository "
             "would measure its agreement with itself.")]
    y = TOP + 0.34
    for code, txt in lims:
        h = 0.60 if code != "L6" else 0.78
        sh = card(s, L, y, 5.92, h, fill=CARD)
        badge = rect(s, L + 0.10, y + 0.10, 0.46, h - 0.20, fill=RED,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        btf = frame_of(badge, 0.02, 0.02)
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(btf, code, size=11.5, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True)
        tf = textbox(s, L + 0.68, y + 0.08, 5.10, h - 0.16)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, txt, size=10.5, color=GREY, after=0, first=True, line=0.95)
        y += h + 0.07
    tfh = textbox(s, 6.86, TOP, 5.85, 0.3)
    para(tfh, "Future work, ordered by value per unit of effort", size=12.5,
         color=INK, bold=True, after=0, first=True)
    fut = [("F1", "**Adjudicate the gold set** — the bottleneck; everything "
            "else is fast. Two annotators, a 500-comment overlap and "
            "Cohen's κ, then 2–3 k comments stratified by bucket and kind.",
            GREEN),
           ("F2", "**Benchmark rather than build** — score the seven heads, "
            "the LLM stance pass and the Stage-1 baseline, reporting "
            "macro-F1 **per language bucket**.", GREEN),
           ("F3", "**Turn the cascade into a measured trade-off** — sweep "
            "the confidence threshold, and plot the **comment-cap lever as "
            "a second curve**.", TEAL),
           ("F4", "**Validate the watchlist separately** — mention detection "
            "(recall matters: a missed alias is invisible) apart from stance "
            "agreement.", TEAL),
           ("F5", "**Learn the routing decision** instead of hand-tuning it "
            "[31] — possible only after F1.", BLUE),
           ("F6", "**Make the image modality real, or delete its weight.** "
            "A weight that has never been anything but zero is not "
            "defensible.", AMBER),
           ("F7", "**Score the agent layer** — groundedness, citation "
            "accuracy, answer relevance and tool-use correctness "
            "[44]–[47].", AMBER)]
    y = TOP + 0.34
    for code, txt, acc in fut:
        h = 0.62 if code in ("F1", "F2") else 0.56
        sh = card(s, 6.86, y, 5.85, h, fill=CARD2)
        badge = rect(s, 6.96, y + 0.10, 0.46, h - 0.20, fill=acc,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        btf = frame_of(badge, 0.02, 0.02)
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(btf, code, size=11.5, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, after=0, first=True)
        tf = textbox(s, 7.54, y + 0.08, 5.03, h - 0.16)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, txt, size=10, color=GREY, after=0, first=True, line=0.95)
        y += h + 0.06
    tfn = textbox(s, L, 6.06, W, 0.3)
    para(tfn, "F1 – F4 would take the project from a demonstrable "
              "engineering result to a defensible empirical one — an "
              "estimated **three to four weeks, dominated almost entirely by "
              "annotation.**", size=11.5, color=INK, after=0, first=True)

    # ------------------------------------------------- 28 references
    s = d.body("References", tag="IEEE style")
    refs_left = [
        "[1]  A. Bhattacharjee et al., “BanglaBERT: Language model "
        "pretraining and benchmarks for low-resource language understanding "
        "evaluation in Bangla,” in *Findings of the ACL: NAACL 2022*, 2022, "
        "pp. 1318–1327.",
        "[2]  K. I. Islam, S. Kar, M. S. Islam, and M. R. Amin, “SentNoB: A "
        "dataset for analysing sentiment on noisy Bangla texts,” in "
        "*Findings of the ACL: EMNLP 2021*, 2021, pp. 3265–3271.",
        "[4]  S. Sazzed, “Abusive content detection in transliterated "
        "Bengali-English social media corpus,” in *Proc. 5th Workshop on "
        "Computational Approaches to Linguistic Code-Switching*, 2021, "
        "pp. 125–130.",
        "[6]  M. A. Nezami and M. N. Huda, “Public sentiment identification "
        "in social media for code-mixed English-Bangla and Banglish text "
        "analysis with machine learning,” in *Proc. IVPAI 2024*, SPIE, 2024.",
        "[9]  F. Alam et al., “A review of Bangla natural language "
        "processing tasks and the utility of transformer models,” "
        "arXiv:2107.03844, 2021.",
        "[10]  P. Patwa et al., “SemEval-2020 Task 9: Overview of sentiment "
        "analysis of code-mixed tweets,” in *Proc. 14th Workshop on Semantic "
        "Evaluation*, 2020, pp. 774–790.",
        "[13]  D. Küçük and F. Can, “Stance detection: A survey,” *ACM "
        "Computing Surveys*, vol. 53, no. 1, pp. 1–37, 2021.",
        "[17]  P. Kralj Novak, J. Smailović, B. Sluban, and I. Mozetič, "
        "“Sentiment of emojis,” *PLOS ONE*, vol. 10, no. 12, e0144296, 2015.",
        "[22]  N. Reimers and I. Gurevych, “Sentence-BERT: Sentence "
        "embeddings using Siamese BERT-networks,” in *Proc. EMNLP-IJCNLP*, "
        "2019, pp. 3980–3990.",
        "[30]  L. Chen, M. Zaharia, and J. Zou, “FrugalGPT: How to use large "
        "language models while reducing cost and improving performance,” "
        "arXiv:2305.05176, 2023.",
    ]
    refs_right = [
        "[31]  I. Ong et al., “RouteLLM: Learning to route LLMs with "
        "preference data,” arXiv:2406.18665, 2024.",
        "[32]  D. Ding et al., “Hybrid LLM: Cost-efficient and quality-aware "
        "query routing,” arXiv:2404.14618, 2024.",
        "[34]  P. Lewis et al., “Retrieval-augmented generation for "
        "knowledge-intensive NLP tasks,” arXiv:2005.11401, 2020.",
        "[36]  S. Yao et al., “ReAct: Synergizing reasoning and acting in "
        "language models,” arXiv:2210.03629, 2022.",
        "[38]  X. Hou, Y. Zhao, S. Wang, and H. Wang, “Model Context "
        "Protocol (MCP): Landscape, security threats, and future research "
        "directions,” arXiv:2503.23278, 2025.",
        "[39]  G. V. Cormack, C. L. A. Clarke, and S. Büttcher, “Reciprocal "
        "rank fusion outperforms Condorcet and individual rank learning "
        "methods,” in *Proc. SIGIR ’09*, 2009, pp. 758–759.",
        "[41]  K. Järvelin and J. Kekäläinen, “Cumulated gain-based "
        "evaluation of IR techniques,” *ACM TOIS*, vol. 20, no. 4, "
        "pp. 422–446, 2002.",
        "[43]  K. Greshake et al., “Not what you've signed up for: "
        "Compromising real-world LLM-integrated applications with indirect "
        "prompt injection,” in *Proc. AISec ’23*, 2023, pp. 79–90.",
        "[44]  L. Zheng et al., “Judging LLM-as-a-judge with MT-Bench and "
        "Chatbot Arena,” arXiv:2306.05685, 2023.",
        "[46]  S. Es, J. James, L. Espinosa-Anke, and S. Schockaert, "
        "“RAGAS: Automated evaluation of retrieval augmented generation,” "
        "arXiv:2309.15217, 2023.",
        "[50]  T. G. Dietterich, “Ensemble methods in machine learning,” in "
        "*Multiple Classifier Systems*, LNCS vol. 1857. Springer, 2000, "
        "pp. 1–15.",
    ]
    for i, col in enumerate((refs_left, refs_right)):
        tf = textbox(s, L + i * 6.17, TOP + 0.02, 5.92, 4.8)
        for j, r in enumerate(col):
            para(tf, r, size=9.5, color=GREY, after=6, line=0.95,
                 first=(j == 0))
    banner(s, L, 5.70, W, 0.60,
           "The full reference list carries **50 sources** in IEEE format; "
           "the twenty-one above are the ones this presentation cites "
           "directly.",
           fill=INK, size=11.5)

    # ------------------------------------------------------ 29 thank you
    s = d.body("", rule=False)
    tf = textbox(s, L, 2.24, W, 1.4)
    para(tf, "Thank you", size=54, color=INK, bold=True,
         align=PP_ALIGN.CENTER, after=6, first=True)
    para(tf, "Questions and discussion", size=18, color=TEAL,
         align=PP_ALIGN.CENTER, after=0)
    rect(s, 5.92, 3.88, 1.5, 0.075, fill=BLUE)
    tf = textbox(s, L, 4.18, W, 0.9)
    para(tf, "Bishwajit Kumar Chakraborty  ·  ID 0242220005101414  ·  "
             "Department of CSE, Daffodil International University", size=13,
         color=GREY, align=PP_ALIGN.CENTER, after=4, first=True)
    para(tf, "Supervised by Dr. Sheak Rashed Haider Noori  ·  Co-supervised "
             "by Dr. Md Alamgir Kabir", size=12, color=MUTED,
         align=PP_ALIGN.CENTER, after=0)
    banner(s, 2.55, 5.22, 8.23, 0.90,
           "“Cheap NLP filters which comments **and** which posts deserve a "
           "large model — and the system proves afterwards, from its own "
           "counters, where the money went and what produced every label.”",
           fill=INK_D, size=12.5, align=PP_ALIGN.CENTER)

    out = d.save()
    print("wrote %s  (%d slides)" % (out, len(d.prs.slides._sldIdLst)))


if __name__ == "__main__":
    build()
